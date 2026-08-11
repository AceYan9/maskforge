import csv
import random
from openpyxl import load_workbook
import tempfile
import shutil
import xlrd

from app.config import settings
from app.utils.minio import get_file


class FileAnalyzer:

    def __init__(self, object_name, file_type):
        self._object_name = object_name
        self._file_type = file_type

    async def __call__(self, *args, **kwargs):
        return await getattr(self, f"_handle_{self._file_type}")()

    async def _handle_csv(self):
        with get_file(self._object_name) as obj:
            samples = []
            total_rows = 0
            total_columns = 0

            with tempfile.NamedTemporaryFile(suffix=".csv", delete=True) as tmp:
                # MinIO流复制到临时文件
                shutil.copyfileobj(obj, tmp)
                # 确保数据写入磁盘
                tmp.flush()

                # 从临时文件读取
                with open(tmp.name, "r", encoding="utf-8", newline="") as f:
                    reader = csv.reader(f)
                    # 第一行作为header
                    headers = next(reader)
                    total_columns = len(headers)
                    for row in reader:
                        sample = dict(zip(headers, row))

                        if len(samples) < settings.FILE_SAMPLE_MAX_SIZE:
                            samples.append(sample)
                        else:
                            index = random.randint(0, total_rows)
                            if index < settings.FILE_SAMPLE_MAX_SIZE:
                                samples[index] = sample
                        total_rows += 1

                return {
                    "rows": total_rows,
                    "columns": total_columns,
                    "headers": headers,
                    "samples": samples,
                }

    async def _handle_xlsx(self):
        with get_file(self._object_name) as obj:
            with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
                shutil.copyfileobj(obj, tmp)
                tmp.flush()
                wb = load_workbook(tmp.name, read_only=True, data_only=True)
                ws = wb.active
                rows = ws.iter_rows()
                headers = [c.value for c in next(rows)]

                samples = []
                total_rows = 0

                for row in rows:
                    values = [c.value for c in row]
                    sample = dict(zip(headers, values))

                    if len(samples) < settings.FILE_SAMPLE_MAX_SIZE:
                        samples.append(sample)
                    else:
                        index = random.randint(0, total_rows)
                        if index < settings.FILE_SAMPLE_MAX_SIZE:
                            samples[index] = sample

                    total_rows += 1

                wb.close()

            return {
                "rows": total_rows,
                "columns": len(headers),
                "headers": headers,
                "samples": samples
            }

    async def _handle_xls(self):
        with get_file(self._object_name) as obj:
            result = {}
            with tempfile.NamedTemporaryFile(suffix=".xls") as tmp:
                shutil.copyfileobj(obj, tmp)
                tmp.flush()
                workbook = xlrd.open_workbook(tmp.name, on_demand=True)

                try:
                    sheet = workbook.sheet_by_index(0)
                    headers = sheet.row_values(0)
                    samples = []
                    for index in range(1, sheet.nrows):
                        row = sheet.row_values(index)
                        sample = dict(zip(headers, row))
                        if len(samples) < settings.FILE_SAMPLE_MAX_SIZE:
                            samples.append(sample)
                        else:
                            random_index = random.randint(0, index - 1)
                            if random_index < settings.FILE_SAMPLE_MAX_SIZE:
                                samples[random_index] = sample

                    result = {
                        "rows": sheet.nrows - 1,
                        "columns": sheet.ncols,
                        "headers": headers,
                        "samples": samples
                    }

                finally:
                    workbook.release_resources()

            return result
